# -*- coding: utf-8 -*-
"""
Created on Thu Jun  5 16:11:05 2025

@author: USUARIO
"""

# Código actualizado con lógica corregida según las indicaciones detalladas
import pandas as pd
import os
import glob
from datetime import datetime, timedelta
import tkinter as tk
from tkinter import filedialog, messagebox
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
def seleccionar_carpeta():
    root = tk.Tk()
    root.withdraw()
    return filedialog.askdirectory(title="Selecciona la carpeta con los archivos TXT")

def unir_txt(carpeta):
    archivos = glob.glob(os.path.join(carpeta, "*.txt"))
    if not archivos:
        print("No se encontraron archivos .txt.")
        return None

    columnas = ["time", "point", "description", "message", "time_unix", "msec", "key2"]
    dfs = []

    for archivo in archivos:
        try:
            df = pd.read_csv(archivo, delimiter='\t', header=None, names=columnas, encoding='latin-1', on_bad_lines='skip')
            dfs.append(df)
        except Exception as e:
            print(f"No se pudo leer {archivo}: {e}")

    if not dfs:
        return None

    df_completo = pd.concat(dfs, ignore_index=True).drop_duplicates()
    df_completo['time'] = pd.to_datetime(df_completo['time'], dayfirst=True, errors='coerce')
    return df_completo.dropna(subset=['time']).sort_values('time')

def pedir_fechas():
    """
    Abre una ventana para que el usuario ingrese fecha de inicio y fecha de fin
    en formato dd/mm/aaaa. Devuelve (fecha_inicio, fecha_fin) como datetime,
    con fecha_inicio a las 00:00:00 y fecha_fin a las 23:59:59.
    Devuelve (None, None) si el usuario cancela.
    """
    resultado = {"inicio": None, "fin": None}

    ventana = tk.Toplevel()
    ventana.title("Rango de fechas a analizar")
    ventana.grab_set()

    tk.Label(ventana, text="Fecha inicio (dd/mm/aaaa):").grid(row=0, column=0, padx=10, pady=10, sticky="e")
    entry_inicio = tk.Entry(ventana)
    entry_inicio.grid(row=0, column=1, padx=10, pady=10)

    tk.Label(ventana, text="Fecha fin (dd/mm/aaaa):").grid(row=1, column=0, padx=10, pady=10, sticky="e")
    entry_fin = tk.Entry(ventana)
    entry_fin.grid(row=1, column=1, padx=10, pady=10)

    def aceptar():
        texto_inicio = entry_inicio.get().strip()
        texto_fin = entry_fin.get().strip()

        try:
            fecha_inicio = datetime.strptime(texto_inicio, "%d/%m/%Y").replace(
                hour=0, minute=0, second=0, microsecond=0
            )
            fecha_fin = datetime.strptime(texto_fin, "%d/%m/%Y").replace(
                hour=23, minute=59, second=59, microsecond=0
            )
        except ValueError:
            messagebox.showerror(
                "Fecha inválida",
                "Una o ambas fechas no tienen el formato correcto (dd/mm/aaaa). Intenta nuevamente."
            )
            return

        if fecha_fin < fecha_inicio:
            messagebox.showerror(
                "Rango inválido",
                "La fecha de fin no puede ser anterior a la fecha de inicio."
            )
            return

        resultado["inicio"] = fecha_inicio
        resultado["fin"] = fecha_fin
        ventana.destroy()

    def cancelar():
        ventana.destroy()

    botones = tk.Frame(ventana)
    botones.grid(row=2, column=0, columnspan=2, pady=10)
    tk.Button(botones, text="Aceptar", command=aceptar).pack(side="left", padx=5)
    tk.Button(botones, text="Cancelar", command=cancelar).pack(side="left", padx=5)

    ventana.wait_window()
    return resultado["inicio"], resultado["fin"]


def filtrar_por_fecha(df, fecha_inicio, fecha_fin):
    """
    Filtra el DataFrame entre fecha_inicio y fecha_fin (inclusive).
    - Si no hay datos en el rango, muestra un mensaje de error y devuelve (None, False).
    - Si hay datos pero no cubren todo el rango solicitado, pregunta al usuario
      si desea continuar con los datos disponibles (botón Sí/No).
    - Si los datos cubren el rango completo, continúa sin preguntar.
    """
    df_filtrado = df[(df['time'] >= fecha_inicio) & (df['time'] <= fecha_fin)].copy()

    if df_filtrado.empty:
        messagebox.showerror(
            "Sin datos",
            f"No se encontraron datos entre {fecha_inicio.strftime('%d/%m/%Y %H:%M:%S')} "
            f"y {fecha_fin.strftime('%d/%m/%Y %H:%M:%S')}."
        )
        return None, False

    min_real = df_filtrado['time'].min()
    max_real = df_filtrado['time'].max()

    tolerancia = timedelta(minutes=1)
    mensajes = []

    if min_real > fecha_inicio + tolerancia:
        mensajes.append(
            f"Los datos disponibles inician el {min_real.strftime('%d/%m/%Y %H:%M:%S')}, "
            f"posterior al inicio solicitado ({fecha_inicio.strftime('%d/%m/%Y %H:%M:%S')})."
        )

    if max_real < fecha_fin - tolerancia:
        mensajes.append(
            f"Los datos disponibles solo llegan hasta el {max_real.strftime('%d/%m/%Y %H:%M:%S')}, "
            f"antes del fin solicitado ({fecha_fin.strftime('%d/%m/%Y %H:%M:%S')})."
        )

    if mensajes:
        continuar = messagebox.askyesno(
            "Datos incompletos",
            "\n".join(mensajes) + "\n\n¿Desea continuar el análisis con los datos disponibles?"
        )
        return df_filtrado, continuar

    return df_filtrado, True


def procesar_alarmas(df):
    df_failvac_alarm = df[
        (df['point'].str[-8:].str.startswith('FAILVCA', na=False)) &
        (df['message'].str.contains('Alarma', case=False, na=False))
    ].copy().dropna(subset=['time']).sort_values('time')

    df_failvac_normal = df[
        (df['point'].str[-8:].str.startswith('FAILVCA', na=False)) &
        (df['message'].str.contains('Normal', case=False, na=False))
    ].copy().dropna(subset=['time']).sort_values('time')

    señales_fco = df[df['point'].str[-8:].str.startswith('F_CO', na=False)]
    señales_fco = señales_fco[señales_fco['message'].str.contains('Alarma|Normal', case=False, na=False)].copy()
    señales_fco = señales_fco.dropna(subset=['time']).sort_values('time')

    resumen_con = {}
    resumen_sin = {}
    detalle_con = []
    detalle_sin = []
    alarmas_activas = {}
    intermitencias_sin = {}

    for _, row in señales_fco.iterrows():
        señal = row['point']
        mensaje = row['message']
        tiempo = row['time']
        base_id = señal[:23]

        if 'Alarma' in mensaje:
            if señal not in alarmas_activas:
                alarmas_activas[señal] = tiempo

        elif 'Normal' in mensaje and señal in alarmas_activas:
            tiempo_inicio = alarmas_activas[señal]
            duracion = tiempo - tiempo_inicio
            segundos = duracion.total_seconds()
            hh_mm_ss = str(duracion)

            # Buscar FAILVAC asociado dentro de las 6 horas antes del F_CO
            ventana_inicio = tiempo_inicio - timedelta(hours=6)
            failvac_match = df_failvac_alarm[
                (df_failvac_alarm['point'].str[:23] == base_id) &
                (df_failvac_alarm['time'] >= ventana_inicio) &
                (df_failvac_alarm['time'] <= tiempo_inicio)
            ]

            if not failvac_match.empty:
                # Tomar el FAILVAC más cercano
                failvac_time = failvac_match['time'].max()
                failvac_point = failvac_match.loc[failvac_match['time'] == failvac_time, 'point'].values[0]

                # Buscar NORMAL posterior al FAILVAC
                failvac_normal_match = df_failvac_normal[
                    (df_failvac_normal['point'] == failvac_point) &
                    (df_failvac_normal['time'] > failvac_time)
                ]
                failvac_normal_time = failvac_normal_match['time'].min() if not failvac_normal_match.empty else None

                # Registrar primero el FAILVAC
                failvac_duracion = (failvac_normal_time - failvac_time) if failvac_normal_time else timedelta(0)
                detalle_con.append({
                    'Señal': failvac_point,
                    'Fecha_Apertura': failvac_time,
                    'Fecha_Cierre': failvac_normal_time,
                    'Duración_HH_MM_SS': str(failvac_duracion),
                    'Duración_Segundos': failvac_duracion.total_seconds()
                })

                # Luego registrar el F_CO con FAILVAC asociado
                detalle_con.append({
                    'Señal': señal,
                    'Fecha_Apertura': tiempo_inicio,
                    'Fecha_Cierre': tiempo,
                    'Duración_HH_MM_SS': hh_mm_ss,
                    'Duración_Segundos': segundos
                })

                resumen_con[señal] = resumen_con.get(señal, timedelta()) + duracion
            else:
                detalle_sin.append({
                    'Señal': señal,
                    'Fecha_Apertura': tiempo_inicio,
                    'Fecha_Cierre': tiempo,
                    'Duración_HH_MM_SS': hh_mm_ss,
                    'Duración_Segundos': segundos
                })
                resumen_sin[señal] = resumen_sin.get(señal, timedelta()) + duracion
                intermitencias_sin[señal] = intermitencias_sin.get(señal, 0) + 1

            del alarmas_activas[señal]

    return resumen_con, resumen_sin, detalle_con, detalle_sin, intermitencias_sin

def guardar_archivos(resumen_con, resumen_sin, detalle_con, detalle_sin,
                     intermitencias_sin, carpeta_salida):

    resumen_df_con = pd.DataFrame([
        [k,
         int(v.total_seconds()//3600),
         int((v.total_seconds()%3600)//60),
         int(v.total_seconds()%60),
         v.total_seconds()]
        for k,v in resumen_con.items()
    ],
    columns=['Señal','Horas','Minutos','Segundos','Total_Segundos'])

    resumen_df_con = resumen_df_con.sort_values(
        "Total_Segundos",
        ascending=False
    )

    resumen_df_sin = pd.DataFrame([
        [k,
         int(v.total_seconds()//3600),
         int((v.total_seconds()%3600)//60),
         int(v.total_seconds()%60),
         v.total_seconds()]
        for k,v in resumen_sin.items()
    ],
    columns=['Señal','Horas','Minutos','Segundos','Total_Segundos'])

    resumen_df_sin = resumen_df_sin.sort_values(
        "Total_Segundos",
        ascending=False
    )

    resumen_df_sin["Intermitencias"] = (
        resumen_df_sin["Señal"]
        .map(intermitencias_sin)
        .fillna(0)
        .astype(int)
    )

    detalle_df_con = pd.DataFrame(detalle_con)
    detalle_df_sin = pd.DataFrame(detalle_sin)

    #=========================================
    # NUEVAS TABLAS
    #=========================================

    detalle_df_sin["Fecha"] = pd.to_datetime(
        detalle_df_sin["Fecha_Apertura"]
    ).dt.date

    intermitencias_fecha = (
        detalle_df_sin
        .groupby("Fecha")
        .size()
        .reset_index(name="Intermitencias")
        .sort_values("Fecha")
    )

    
    #=========================================
    # RESUMEN DIARIO POR EQUIPO (NUEVO)
    #=========================================
    detalle_df_sin["Fecha"] = pd.to_datetime(detalle_df_sin["Fecha_Apertura"]).dt.date
    detalle_df_sin["Duracion_td"] = pd.to_timedelta(detalle_df_sin["Duración_HH_MM_SS"])

    filas=[]
    for fecha,g in detalle_df_sin.groupby("Fecha"):
        total=len(g)
        neq=g["Señal"].nunique()
        filas.append({
            "Fecha":fecha,
            "Total Intermitencias":total,
            "Nº Equipos":neq,
            "Equipo":"",
            "Intermitencias":"",
            "Tiempo acumulado":""
        })
        gp=(g.groupby("Señal")
              .agg(Intermitencias=("Señal","size"),
                   Tiempo=("Duracion_td","sum"))
              .reset_index())
        gp=gp.sort_values("Intermitencias",ascending=False)
        for _,r in gp.iterrows():
            filas.append({
                "Fecha":"",
                "Total Intermitencias":"",
                "Nº Equipos":"",
                "Equipo":r["Señal"],
                "Intermitencias":int(r["Intermitencias"]),
                "Tiempo acumulado":str(r["Tiempo"])
            })
    resumen_diario_equipos=pd.DataFrame(filas)

    top_equipos = (
        resumen_df_sin[["Señal","Intermitencias"]]
        .sort_values(
            "Intermitencias",
            ascending=False
        )
    )

    resumen_path = os.path.join(
        carpeta_salida,
        "Resumen_Tiempos_Desconexion.xlsx"
    )

    detalle_path = os.path.join(
        carpeta_salida,
        "Detalle_Tiempos_Desconexion.xlsx"
    )

    with pd.ExcelWriter(resumen_path,
                        engine="openpyxl") as writer:

        resumen_df_sin.to_excel(
            writer,
            sheet_name="Sin_FAILVAC",
            index=False
        )

        resumen_df_con.to_excel(
            writer,
            sheet_name="Asociadas_FAILVAC",
            index=False
        )

        intermitencias_fecha.to_excel(
            writer,
            sheet_name="Intermitencias_por_Fecha",
            index=False
        )

        resumen_diario_equipos.to_excel(writer,sheet_name="Resumen_Diario_Equipos",index=False)

        top_equipos.to_excel(
            writer,
            sheet_name="Top_Equipos",
            index=False
        )

    with pd.ExcelWriter(detalle_path,
                        engine="openpyxl") as writer:

        detalle_df_sin.to_excel(
            writer,
            sheet_name="Sin_FAILVAC",
            index=False
        )

        detalle_df_con.to_excel(
            writer,
            sheet_name="Asociadas_FAILVAC",
            index=False
        )

    #=========================================
    # GRAFICOS
    #=========================================

    wb = load_workbook(resumen_path)

    ws_top = wb["Top_Equipos"]

    ws_fecha = wb["Intermitencias_por_Fecha"]

    ws_graf = wb.create_sheet("Graficos")

    #-----------------------------
    # Grafico Equipos
    #-----------------------------

    chart1 = BarChart()

    datos = Reference(
        ws_top,
        min_col=2,
        min_row=1,
        max_row=ws_top.max_row
    )

    categorias = Reference(
        ws_top,
        min_col=1,
        min_row=2,
        max_row=ws_top.max_row
    )

    chart1.add_data(
        datos,
        titles_from_data=True
    )

    chart1.set_categories(categorias)

    chart1.title = "Intermitencias por Equipo"

    chart1.height = 10
    chart1.width = 18

    ws_graf.add_chart(chart1,"A1")

    #-----------------------------
    # Grafico por Fecha
    #-----------------------------

    chart2 = BarChart()

    datos2 = Reference(
        ws_fecha,
        min_col=2,
        min_row=1,
        max_row=ws_fecha.max_row
    )

    categorias2 = Reference(
        ws_fecha,
        min_col=1,
        min_row=2,
        max_row=ws_fecha.max_row
    )

    chart2.add_data(
        datos2,
        titles_from_data=True
    )

    chart2.set_categories(categorias2)

    chart2.title = "Intermitencias por Fecha"

    chart2.height = 10
    chart2.width = 18

    ws_graf.add_chart(chart2,"A22")

    wb.save(resumen_path)

    print("Archivos generados correctamente.")

# === EJECUCIÓN PRINCIPAL ===
root = tk.Tk()
root.withdraw()

carpeta = seleccionar_carpeta()
if carpeta:
    df = unir_txt(carpeta)
    if df is not None:
        fecha_inicio, fecha_fin = pedir_fechas()

        if fecha_inicio is None or fecha_fin is None:
            print("No se ingresó un rango de fechas. Proceso cancelado.")
        else:
            df_filtrado, continuar = filtrar_por_fecha(df, fecha_inicio, fecha_fin)

            if df_filtrado is not None and continuar:
                r_con, r_sin, d_con, d_sin, intermitencias = procesar_alarmas(df_filtrado)
                guardar_archivos(r_con, r_sin, d_con, d_sin, intermitencias, carpeta)
            elif df_filtrado is not None and not continuar:
                print("Proceso cancelado por el usuario debido a datos incompletos.")
            else:
                print("No se encontraron datos para el rango de fechas solicitado.")
    else:
        print("No se pudo procesar el DataFrame.")
else:
    print("No se seleccionó carpeta.")

root.destroy()